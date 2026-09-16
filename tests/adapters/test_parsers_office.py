"""Tests for the docx/xlsx/pptx office document parser adapters."""

import pytest

from porsuk.adapters.parsers.docx import DocxParser
from porsuk.adapters.parsers.pptx import PptxParser
from porsuk.adapters.parsers.xlsx import XlsxParser
from porsuk.core.ports import Parser


@pytest.mark.parametrize("cls", [DocxParser, XlsxParser, PptxParser])
def test_conform_to_the_parser_protocol(cls):
    assert isinstance(cls(), Parser)


def test_module_names_do_not_shadow_the_libraries_they_import():
    """docx.py and pptx.py are named after their registry keys, matching the
    fake.py / inmemory.py convention. Absolute imports make that safe, but the
    failure mode would be silent, so it is asserted rather than assumed."""
    import porsuk.adapters.parsers.docx as our_docx
    import porsuk.adapters.parsers.pptx as our_pptx

    assert our_docx.__name__ == "porsuk.adapters.parsers.docx"
    assert our_pptx.__name__ == "porsuk.adapters.parsers.pptx"


@pytest.mark.parametrize(
    ("cls", "good", "bad"),
    [
        (DocxParser, "/tmp/a.docx", "/tmp/a.pdf"),
        (XlsxParser, "/tmp/a.xlsx", "/tmp/a.docx"),
        (PptxParser, "/tmp/a.pptx", "/tmp/a.xlsx"),
    ],
)
def test_each_handles_only_its_own_extension(cls, good, bad):
    parser = cls()
    assert parser.can_handle(good)
    assert parser.can_handle(good.upper())
    assert not parser.can_handle(bad)


def test_docx_reads_style_headings_as_an_outline(fixtures_dir):
    """Tier 1: Word style headings are real structure."""
    outcome = DocxParser().parse(str(fixtures_dir / "docx" / "styled_headings.docx"))
    assert outcome.status == "ok"
    outline = outcome.document.outline
    assert outline.source == "toc"
    titles = [n.title for n in outline.nodes]
    assert "Tedarik Sözleşmesi" in titles
    nested = [c.title for n in outline.nodes for c in n.children]
    assert "1. Taraflar" in nested
    assert "2. Mali Hükümler" in nested


def test_docx_body_text_becomes_blocks(fixtures_dir):
    outcome = DocxParser().parse(str(fixtures_dir / "docx" / "styled_headings.docx"))
    text = "\n".join(b.text for b in outcome.document.blocks)
    assert "sözleşme" in text
    assert any(b.kind == "heading" for b in outcome.document.blocks)


def test_docx_reports_one_page(fixtures_dir):
    """Word has no page count without rendering. One page means chars_per_page
    is the whole document, which is harmless: a docx is never an OCR
    candidate."""
    outcome = DocxParser().parse(str(fixtures_dir / "docx" / "styled_headings.docx"))
    assert outcome.document.page_count == 1


def test_xlsx_counts_sheets_as_pages(fixtures_dir):
    """openpyxl, sheet by sheet."""
    outcome = XlsxParser().parse(str(fixtures_dir / "xlsx" / "two_sheets.xlsx"))
    assert outcome.status == "ok"
    assert outcome.document.page_count == 2


def test_xlsx_emits_one_table_block_per_populated_sheet(fixtures_dir):
    outcome = XlsxParser().parse(str(fixtures_dir / "xlsx" / "two_sheets.xlsx"))
    tables = [b for b in outcome.document.blocks if b.kind == "table"]
    assert len(tables) == 2
    assert "F-001" in "\n".join(b.text for b in tables)


def test_xlsx_sheet_names_become_the_outline(fixtures_dir):
    outcome = XlsxParser().parse(str(fixtures_dir / "xlsx" / "two_sheets.xlsx"))
    assert outcome.document.outline.source == "toc"
    assert [n.title for n in outcome.document.outline.nodes] == ["Ödemeler", "Teslimat"]


def test_xlsx_emits_a_heading_block_before_each_sheets_table(fixtures_dir):
    """docx.py and pptx.py both emit a heading block for every outline node;
    xlsx.py did not, so tier 1 (`sections_from_outline`) had nothing whose
    text started with the sheet name to match against and every spreadsheet
    chunk fell through to the semantic tier with `section_path=None`."""
    outcome = XlsxParser().parse(str(fixtures_dir / "xlsx" / "two_sheets.xlsx"))
    blocks = outcome.document.blocks
    assert [b.kind for b in blocks] == ["heading", "table", "heading", "table"]
    assert blocks[0].text == "Ödemeler"
    assert blocks[0].page_no == 1
    assert blocks[2].text == "Teslimat"
    assert blocks[2].page_no == 2


def test_xlsx_sections_from_outline_matches_the_sheet_headings(fixtures_dir):
    """The fix must actually close the tier-1 gap it targets, not just add a
    block that happens to carry the right text."""
    from porsuk.ingestion.chunking.toc import sections_from_outline

    outcome = XlsxParser().parse(str(fixtures_dir / "xlsx" / "two_sheets.xlsx"))
    sections = sections_from_outline(outcome.document)
    assert sections is not None
    assert [s.path for s in sections] == ["Ödemeler", "Teslimat"]


def test_xlsx_chunks_carry_the_sheet_name_as_section_path(fixtures_dir):
    from porsuk.core.config import ChunkingConfig
    from porsuk.ingestion.chunking import chunk

    outcome = XlsxParser().parse(str(fixtures_dir / "xlsx" / "two_sheets.xlsx"))
    chunks = chunk(outcome.document, "doc1", ChunkingConfig())
    assert chunks
    assert all(c.section_path in {"Ödemeler", "Teslimat"} for c in chunks)


def test_pptx_counts_slides_as_pages(fixtures_dir):
    outcome = PptxParser().parse(str(fixtures_dir / "pptx" / "two_slides.pptx"))
    assert outcome.status == "ok"
    assert outcome.document.page_count == 2


def test_pptx_slide_titles_become_the_outline(fixtures_dir):
    """Tier 1 explicitly names pptx slide titles."""
    outcome = PptxParser().parse(str(fixtures_dir / "pptx" / "two_slides.pptx"))
    assert outcome.document.outline.source == "toc"
    assert [n.title for n in outcome.document.outline.nodes] == [
        "Tedarik Süreci",
        "Supply Process",
    ]


def test_pptx_blocks_carry_their_slide_number(fixtures_dir):
    outcome = PptxParser().parse(str(fixtures_dir / "pptx" / "two_slides.pptx"))
    assert {b.page_no for b in outcome.document.blocks} == {1, 2}


@pytest.mark.parametrize(
    ("cls", "name"), [(DocxParser, "a.docx"), (XlsxParser, "a.xlsx"), (PptxParser, "a.pptx")]
)
def test_a_corrupt_file_fails_without_raising(cls, tmp_path, name):
    """A Result, not an exception."""
    path = tmp_path / name
    path.write_bytes(b"not an office file at all")
    outcome = cls().parse(str(path))
    assert outcome.status == "failed"
    assert outcome.document is None
    assert outcome.error


def test_docx_table_cells_become_a_table_block(fixtures_dir):
    """Word keeps amounts, dates and terms in tables, and `Document.paragraphs`
    never descends into one. Dropping them is invisible downstream: with no
    table block emitted, `empty_table_blocks` stays 0 and the gate reports a
    clean parse over content that was never extracted."""
    outcome = DocxParser().parse(str(fixtures_dir / "docx" / "with_table.docx"))
    blocks = outcome.document.blocks
    assert any(b.kind == "table" for b in blocks)
    table = next(b for b in blocks if b.kind == "table")
    assert "F-001" in table.text
    assert "Nakliye bedeli" in table.text
    assert "12.500 TL" in table.text


def test_docx_keeps_body_and_table_in_document_order(fixtures_dir):
    outcome = DocxParser().parse(str(fixtures_dir / "docx" / "with_table.docx"))
    kinds = [b.kind for b in outcome.document.blocks]
    assert kinds.index("table") > kinds.index("heading")
    trailing = [b for b in outcome.document.blocks if "30 gün" in b.text]
    assert trailing, "the paragraph after the table must survive"


def test_pptx_table_cells_become_a_table_block(fixtures_dir):
    """A pptx table is a GraphicFrame and has no text frame at all."""
    outcome = PptxParser().parse(str(fixtures_dir / "pptx" / "with_table.pptx"))
    blocks = outcome.document.blocks
    assert any(b.kind == "table" for b in blocks)
    table = next(b for b in blocks if b.kind == "table")
    assert "Nakliye" in table.text
    assert "12.500 TL" in table.text


def test_pptx_emits_each_slide_title_exactly_once(fixtures_dir):
    """`slide.shapes.title` builds a fresh proxy on every access, so an
    identity test against it never matches a shape yielded by iteration and
    the title is emitted twice: once as a heading, once as body text."""
    outcome = PptxParser().parse(str(fixtures_dir / "pptx" / "two_slides.pptx"))
    texts = [b.text for b in outcome.document.blocks]
    assert texts.count("Tedarik Süreci") == 1
    assert texts.count("Supply Process") == 1


def test_a_workbook_with_blank_trailing_sheets_is_not_called_scanned(tmp_path):
    """`page_count` counted every sheet while blocks were emitted only for
    populated ones, so blank trailing sheets, entirely normal in a real
    workbook, diluted chars_per_page until the gate concluded 'likely
    scanned' and asked for OCR on a spreadsheet."""
    from openpyxl import Workbook

    from porsuk.core.config import QualityConfig
    from porsuk.ingestion.quality import decide, measure

    book = Workbook()
    sheet = book.active
    sheet.title = "Ödemeler"
    sheet.append(["Fatura No", "Tutar", "Vade"])
    sheet.append(["F-001", 15000, "2024-04-15"])
    for i in range(9):
        book.create_sheet(f"Bos{i}")
    path = tmp_path / "trailing.xlsx"
    book.save(path)

    outcome = XlsxParser().parse(str(path))
    gate = decide(measure(outcome.document), QualityConfig())
    assert gate.escalate_to is None, f"a spreadsheet was sent to OCR: {gate.reason}"
