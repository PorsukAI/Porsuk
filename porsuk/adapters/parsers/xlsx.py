"""Excel parser, one table block per sheet ('sheet başına').

Each populated sheet also gets a heading block carrying its sheet name,
emitted immediately before that sheet's table block, the same convention
docx.py and pptx.py use for their outline nodes. Tier 1 (`sections_from_
outline`) matches an outline entry to the first block whose text
starts with the entry's title; without a heading block, a spreadsheet's
outline entries had nothing to match and every spreadsheet chunk fell
through to the semantic tier with `section_path=None`.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from porsuk.core.models import Block, Outline, OutlineNode, ParsedDocument, ParseOutcome
from porsuk.core.registry import register

_NAME = "xlsx"


def _failed(error: str) -> ParseOutcome:
    return ParseOutcome(document=None, status="failed", quality=0.0, parser_used=_NAME, error=error)


def _sheet_text(sheet) -> str:
    """Rows as pipe-separated lines, the same shape tables are kept in elsewhere."""
    lines = []
    for row in sheet.iter_rows(values_only=True):
        cells = ["" if c is None else str(c) for c in row]
        if any(cell.strip() for cell in cells):
            lines.append(" | ".join(cells))
    return "\n".join(lines)


@register("parser", _NAME)
class XlsxParser:
    name = _NAME
    cost = 10

    def can_handle(self, path: str) -> bool:
        return Path(path).suffix.lower() == ".xlsx"

    def parse(self, path: str) -> ParseOutcome:
        try:
            document = self._parse(path)
        except Exception as exc:  # noqa: BLE001
            return _failed(f"{type(exc).__name__}: {exc}")
        return ParseOutcome(
            document=document, status="ok", quality=0.0, parser_used=_NAME, error=None
        )

    def _parse(self, path: str) -> ParsedDocument:
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            blocks: list[Block] = []
            nodes: list[OutlineNode] = []
            populated = 0
            for index, sheet in enumerate(book.worksheets, start=1):
                nodes.append(OutlineNode(title=sheet.title, level=1, page_no=index))
                text = _sheet_text(sheet)
                # An entirely empty sheet emits nothing rather than an empty
                # table block: it is a normal thing for a workbook to contain,
                # and emitting one would trip the Docling trigger. The
                # heading block follows the same rule and is skipped too - a
                # heading with no table under it would be a bare sheet-name
                # chunk carrying no information beyond what the outline
                # already has, for a workbook with nothing wrong with it.
                if text:
                    populated += 1
                    blocks.append(Block(text=sheet.title, page_no=index, kind="heading", bold=True))
                    blocks.append(Block(text=text, page_no=index, kind="table"))
            return ParsedDocument(
                path=path,
                doc_type="xlsx",
                blocks=tuple(blocks),
                outline=Outline(nodes=tuple(nodes), source="toc")
                if nodes
                else Outline(source="none"),
                # Populated sheets, not every sheet. Blocks are emitted only
                # for sheets with content, so counting blank trailing ones,
                # entirely normal in a real workbook, diluted chars_per_page
                # until the gate concluded "likely scanned" and asked for OCR
                # on a spreadsheet. Floored at 1 so an empty workbook does not
                # divide by zero.
                page_count=max(1, populated),
            )
        finally:
            book.close()
