"""Word parser.

Named after its registry key, matching the fake.py / inmemory.py convention.
The name shadows the `docx` package it imports, which is safe: Python 3
imports are absolute, so `from docx import Document` still finds the installed
library.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from porsuk.core.models import Block, Outline, OutlineNode, ParsedDocument, ParseOutcome
from porsuk.core.registry import register

_NAME = "docx"


def _failed(error: str) -> ParseOutcome:
    return ParseOutcome(document=None, status="failed", quality=0.0, parser_used=_NAME, error=error)


def _heading_level(style_name: str) -> int | None:
    """python-docx exposes headings as style names: 'Heading 1', 'Title'."""
    if style_name == "Title":
        return 1
    if style_name.startswith("Heading "):
        tail = style_name.removeprefix("Heading ").strip()
        return int(tail) + 1 if tail.isdigit() else None
    return None


def _table_text(table: Table) -> str:
    """Rows as pipe-separated lines: the shape xlsx.py already emits and the
    shape `table_to_markdown` renders from."""
    lines = []
    for row in table.rows:
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        if any(cells):
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def _iter_body(source) -> list[Paragraph | Table]:
    """Paragraphs and tables in document order.

    `Document.paragraphs` returns only top-level body paragraphs and never
    descends into a table, so iterating it drops every table cell in the file.
    Walking the body element is python-docx's documented way to get both, and
    order matters because the chunker attributes blocks to the heading above
    them.
    """
    items: list[Paragraph | Table] = []
    for child in source.element.body.iterchildren():
        if child.tag == qn("w:p"):
            items.append(Paragraph(child, source))
        elif child.tag == qn("w:tbl"):
            items.append(Table(child, source))
    return items


def _nest(entries: list[tuple[str, int]]) -> tuple[OutlineNode, ...]:
    roots: list[dict] = []
    stack: list[dict] = []
    for title, level in entries:
        node = {"title": title, "level": level, "children": []}
        while stack and stack[-1]["level"] >= level:
            stack.pop()
        if stack:
            stack[-1]["children"].append(node)
        else:
            roots.append(node)
        stack.append(node)

    def build(raw: dict) -> OutlineNode:
        return OutlineNode(
            title=raw["title"],
            level=raw["level"],
            page_no=None,
            children=tuple(build(c) for c in raw["children"]),
        )

    return tuple(build(r) for r in roots)


@register("parser", _NAME)
class DocxParser:
    name = _NAME
    cost = 10

    def can_handle(self, path: str) -> bool:
        return Path(path).suffix.lower() == ".docx"

    def parse(self, path: str) -> ParseOutcome:
        try:
            document = self._parse(path)
        except Exception as exc:  # noqa: BLE001
            return _failed(f"{type(exc).__name__}: {exc}")
        return ParseOutcome(
            document=document, status="ok", quality=0.0, parser_used=_NAME, error=None
        )

    def _parse(self, path: str) -> ParsedDocument:
        source = Document(path)
        blocks: list[Block] = []
        headings: list[tuple[str, int]] = []
        for item in _iter_body(source):
            if isinstance(item, Table):
                text = _table_text(item)
                if text:
                    blocks.append(Block(text=text, page_no=1, kind="table"))
                continue
            text = item.text.strip()
            if not text:
                continue
            level = _heading_level(item.style.name if item.style else "")
            if level is None:
                blocks.append(Block(text=text, page_no=1, kind="text"))
            else:
                headings.append((text, level))
                blocks.append(Block(text=text, page_no=1, kind="heading", bold=True))
        outline = (
            Outline(nodes=_nest(headings), source="toc") if headings else Outline(source="none")
        )
        return ParsedDocument(
            path=path,
            doc_type="docx",
            blocks=tuple(blocks),
            outline=outline,
            # Word has no page count without rendering. One page makes
            # chars_per_page the whole document, which is harmless: the
            # OCR branch is about scanned PDFs, and a docx never reaches it.
            page_count=1,
        )
