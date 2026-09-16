"""PowerPoint parser, one page per slide ('slayt başına').

Named after its registry key; shadows the `pptx` package it imports, which is
safe under Python 3 absolute imports.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from porsuk.core.models import Block, Outline, OutlineNode, ParsedDocument, ParseOutcome
from porsuk.core.registry import register

_NAME = "pptx"


def _table_text(table) -> str:
    """Rows as pipe-separated lines, matching docx.py and xlsx.py."""
    lines = []
    for row in table.rows:
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        if any(cells):
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def _iter_shapes(shapes):
    """Flatten groups: a shape inside a GroupShape is never yielded by a plain
    iteration, so text grouped on a slide would be dropped."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _failed(error: str) -> ParseOutcome:
    return ParseOutcome(document=None, status="failed", quality=0.0, parser_used=_NAME, error=error)


@register("parser", _NAME)
class PptxParser:
    name = _NAME
    cost = 10

    def can_handle(self, path: str) -> bool:
        return Path(path).suffix.lower() == ".pptx"

    def parse(self, path: str) -> ParseOutcome:
        try:
            document = self._parse(path)
        except Exception as exc:  # noqa: BLE001
            return _failed(f"{type(exc).__name__}: {exc}")
        return ParseOutcome(
            document=document, status="ok", quality=0.0, parser_used=_NAME, error=None
        )

    def _parse(self, path: str) -> ParsedDocument:
        presentation = Presentation(path)
        blocks: list[Block] = []
        nodes: list[OutlineNode] = []
        slides = list(presentation.slides)
        for index, slide in enumerate(slides, start=1):
            # Captured once. `shapes.title` builds a fresh proxy on every
            # access, so `shape is slide.shapes.title` never matches a shape
            # yielded by iteration and every title was emitted twice.
            title_shape = slide.shapes.title
            title_element = title_shape._element if title_shape is not None else None
            title = title_shape.text.strip() or None if title_shape is not None else None
            if title:
                nodes.append(OutlineNode(title=title, level=1, page_no=index))
                blocks.append(Block(text=title, page_no=index, kind="heading", bold=True))
            for shape in _iter_shapes(slide.shapes):
                if shape._element is title_element:
                    continue
                if shape.has_table:
                    text = _table_text(shape.table)
                    if text:
                        blocks.append(Block(text=text, page_no=index, kind="table"))
                    continue
                if not shape.has_text_frame:
                    continue
                text = shape.text_frame.text.strip()
                if text:
                    blocks.append(Block(text=text, page_no=index, kind="text"))
        return ParsedDocument(
            path=path,
            doc_type="pptx",
            blocks=tuple(blocks),
            outline=Outline(nodes=tuple(nodes), source="toc") if nodes else Outline(source="none"),
            page_count=len(slides),
        )
